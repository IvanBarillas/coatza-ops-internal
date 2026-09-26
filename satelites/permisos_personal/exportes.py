"""Archivos de Excel de las solicitudes y del calendario de ausencias (openpyxl, en memoria)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

CABECERA = PatternFill("solid", fgColor="1F2937")


def _libro(titulo, columnas, filas, anchos):
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]
    hoja.append(columnas)
    for celda in hoja[1]:
        celda.font, celda.fill, celda.alignment = Font(bold=True, color="FFFFFF"), CABECERA, Alignment(vertical="center")
    for fila in filas:
        hoja.append(fila)
    for letra, ancho in zip("ABCDEFGHIJKL", anchos):
        hoja.column_dimensions[letra].width = ancho
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _texto(valor):
    """Evita que una celda de texto libre se interprete como fórmula al abrirla en Excel."""
    valor = "" if valor is None else str(valor)
    return "'" + valor if valor[:1] in ("=", "+", "-", "@") else valor


def solicitudes_xlsx(solicitudes):
    columnas = ["Empleado", "Tipo de personal", "Sede", "Solicitud", "Del", "Al", "Días", "Estado", "Validada", "Validó", "Bajo el mínimo", "Comentarios"]
    filas = [
        [
            _texto(s.empleado.nombre), s.empleado.get_tipo_display(), _texto(s.sede_nombre), s.get_tipo_display(), s.fecha_inicio, s.fecha_fin, s.dias,
            s.get_estatus_display(), "Sí" if s.validado else "No", _texto(getattr(s.validado_por, "full_name", "") if s.validado_por else ""),
            "Sí" if s.bajo_umbral else "No", _texto(s.comentarios),
        ]
        for s in solicitudes
    ]
    return _libro("Solicitudes", columnas, filas, [28, 16, 24, 16, 12, 12, 7, 12, 10, 24, 14, 40])


def ausencias_xlsx(semanas, mes):
    """Una fila por persona ausente y día del mes (solo días del mes, sin repetir los de la semana vecina)."""
    columnas = ["Fecha", "Empleado", "Sede", "Solicitud", "Validada", "Sede bajo el mínimo ese día"]
    filas = []
    for semana in semanas:
        for dia in semana:
            if dia["fecha"].month != mes:
                continue
            bajo = ", ".join(b["sede"] for b in dia["bajo"])
            for f in dia["fuera"]:
                filas.append([dia["fecha"], _texto(f["nombre"]), _texto(f["sede"]), "Día económico" if f["tipo"] == "economico" else "Vacaciones", "Sí" if f["validado"] else "No", bajo])
    return _libro("Ausencias", columnas, filas, [12, 28, 24, 16, 10, 30])
