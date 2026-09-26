import datetime
import io

from django.core import mail
from django.urls import reverse
from openpyxl import load_workbook

from apps.security.models import UserAppRole
from satelites.permisos_personal import services as svc
from satelites.permisos_personal.models import UmbralSede
from satelites.permisos_personal.permissions import PermisosPersonalPermissions as P

from .base import LUNES, BasePermisos, d


def url(nombre, *args):
    return reverse(f"permisos_personal:{nombre}", args=args)


def hoja(respuesta):
    return list(load_workbook(io.BytesIO(respuesta.content)).active.iter_rows(values_only=True))


class ExportesTests(BasePermisos):
    def test_excel_de_solicitudes_respeta_los_filtros_y_los_permisos(self):
        svc.crear_solicitud(self.ana, usuario=self.ana.usuario, tipo=svc.VACACIONES, fecha_inicio=LUNES, fecha_fin=LUNES + datetime.timedelta(days=2), comentarios="=SUMA(1;1)")
        svc.crear_solicitud(self.cris, usuario=self.cris.usuario, tipo=svc.ECONOMICO, fecha_inicio=LUNES, fecha_fin=LUNES)
        self.client.force_login(self.control)
        r = self.client.get(url("solicitudes_exportar"), {"anio": 2026, "estado": "todas"})
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn("solicitudes-2026-todas.xlsx", r["Content-Disposition"])
        filas = hoja(r)
        self.assertEqual(filas[0][:4], ("Empleado", "Tipo de personal", "Sede", "Solicitud"))
        self.assertEqual({f[0] for f in filas[1:]}, {"Ana Sindicalizada", "Cris Sindicalizado"})
        ana = next(f for f in filas[1:] if f[0] == "Ana Sindicalizada")
        self.assertEqual((ana[3], ana[6], ana[8]), ("Vacaciones", 3, "No"))
        self.assertTrue(ana[11].startswith("'="))        # el texto libre no se ejecuta como fórmula
        solo = hoja(self.client.get(url("solicitudes_exportar"), {"anio": 2026, "estado": "todas", "tipo": "economico"}))
        self.assertEqual([f[0] for f in solo[1:]], ["Cris Sindicalizado"])
        self.client.force_login(self.ana.usuario)
        self.assertIn(self.client.get(url("solicitudes_exportar")).status_code, (302, 403))

    def test_excel_del_calendario_lista_ausencias_del_mes(self):
        svc.crear_solicitud(self.ana, usuario=self.ana.usuario, tipo=svc.VACACIONES, fecha_inicio=LUNES, fecha_fin=LUNES + datetime.timedelta(days=1))
        self.client.force_login(self.control)
        r = self.client.get(url("calendario_exportar"), {"anio": 2026, "mes": 3})
        filas = hoja(r)
        self.assertEqual(filas[0][:3], ("Fecha", "Empleado", "Sede"))
        self.assertEqual([(f[1], f[3]) for f in filas[1:]], [("Ana Sindicalizada", "Vacaciones")] * 2)
        self.assertEqual(len(hoja(self.client.get(url("calendario_exportar"), {"anio": 2026, "mes": 4}))), 1)  # solo encabezado
        self.assertEqual(self.client.get(url("calendario_exportar"), {"anio": "x", "mes": "99"}).status_code, 200)


class AvisosCoberturaTests(BasePermisos):
    def setUp(self):
        UmbralSede.objects.create(sede_uuid=self.sede_a.pk, sede_nombre="Palacio", minimo=2)   # Palacio: Ana y Beto

    def correos(self):
        return [(m.to[0], m.subject, m.body) for m in mail.outbox]

    def test_una_solicitud_bajo_el_minimo_avisa_a_control_y_no_al_solicitante(self):
        with self.captureOnCommitCallbacks(execute=True):
            svc.crear_solicitud(self.ana, usuario=self.ana.usuario, tipo=svc.VACACIONES, fecha_inicio=LUNES, fecha_fin=LUNES + datetime.timedelta(days=1))
        (para, asunto, cuerpo), = self.correos()
        self.assertEqual(para, "control@example.test")
        self.assertEqual(asunto, "Cobertura baja en Palacio: Ana Sindicalizada")
        self.assertIn("02/03/2026: 1 presente(s), mínimo 2", cuerpo)
        self.assertIn("no se bloquea", cuerpo)

    def test_no_avisa_si_la_cobertura_alcanza_o_con_avisar_falso(self):
        UmbralSede.objects.filter(sede_uuid=self.sede_a.pk).update(minimo=1)               # Palacio: con uno presente alcanza
        with self.captureOnCommitCallbacks(execute=True):
            svc.crear_solicitud(self.ana, usuario=self.ana.usuario, tipo=svc.VACACIONES, fecha_inicio=LUNES, fecha_fin=LUNES)
            UmbralSede.objects.filter(sede_uuid=self.sede_a.pk).update(minimo=2)
            svc.crear_solicitud(self.beto, usuario=self.beto.usuario, tipo=svc.VACACIONES, fecha_inicio=d(3, 9), fecha_fin=d(3, 9), avisar=False)
        self.assertEqual(self.correos(), [])

    def test_avisa_a_todos_los_de_control_menos_al_que_solicita(self):
        otro = self.usuario("owner@example.test", "owner", "Olga", "Owner")
        with self.captureOnCommitCallbacks(execute=True):
            svc.crear_solicitud(self.ana, usuario=self.control, tipo=svc.VACACIONES, fecha_inicio=LUNES, fecha_fin=LUNES)   # la registra control por ella
        self.assertEqual(sorted(p for p, _, _ in self.correos()), ["owner@example.test"])
